"""Contract resources must survive real file rendering with their conditions intact."""

import base64
import io
import re

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pypdf import PdfReader

from app.contexts.conversations.application.agent_skills import read_skill
from app.contexts.conversations.application.agent_tools import FileCreateInput
from app.contexts.documents.infrastructure.document_files import operate


TEMPLATES = {
    "tea-purchase": ("隐蔽瑕疵", "未约定定金时", "不直接转为该工厂的合同义务"),
    "tea-export": ("所有权转移", "未签约的工厂", "实际营业地"),
    "export-agency": ("主管机关审核确定", "不当然转移法定责任", "自己／委托方"),
    "tea-distribution": ("适用的竞争规则", "经销方与进口商分别识别", "库存回购"),
}


@pytest.mark.parametrize("template", TEMPLATES)
@pytest.mark.parametrize("format", ["docx", "pdf"])
def test_contract_templates_render_with_signatures_and_material_conditions(template, format):
    content = read_skill("draft", f"assets/contracts/{template}.md")
    result = operate("create", {"format": format, "layout": "contract", "content": content})
    assert result["validation"]["reopened"]
    assert result["validation"]["layout"] == "contract"
    data = base64.b64decode(result["file_base64"])
    if format == "docx":
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        )
        assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
        assert round(doc.sections[0].page_width.mm) == 210
        assert round(doc.sections[0].left_margin.mm) == 25
        assert doc.styles["Normal"].font.size.pt == 12
        assert str(doc.styles["Heading 2"].font.color.rgb) == "000000"
        assert "PAGE" in doc.sections[0].footer._element.xml
        assert doc.tables[-1].rows[0]._tr.find("w:trPr/w:tblHeader", doc._element.nsmap) is not None
        assert doc.tables[-1].rows[1]._tr.find("w:trPr/w:cantSplit", doc._element.nsmap) is not None
        assert doc.paragraphs[0].style.element.get(qn("w:styleId")) == "Title"
    else:
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() for page in reader.pages)
        assert 2 <= len(reader.pages) <= 8
        assert "签署栏" in reader.pages[-1].extract_text()
        assert "本栏对应《" in reader.pages[-1].extract_text()
        for index, page in enumerate(reader.pages):
            assert f"第 {index + 1} 页" in page.extract_text()
            assert abs(float(page.mediabox.width) - 595.276) < 1
    normalized = re.sub(r"\s+", "", text)
    for phrase in (*TEMPLATES[template], "签署栏", "[待补充]"):
        assert re.sub(r"\s+", "", phrase) in normalized
    assert "GF—2025—0151" not in text
    assert not any(word in content for word in ("不默认", "不自动加入", "凭模板", "Agent", "上游"))


def test_contract_layout_cannot_be_misapplied_to_another_format():
    with pytest.raises(ValueError, match="仅支持Word和PDF"):
        operate("create", {"format": "txt", "layout": "contract", "content": "合同正文"})


@pytest.mark.parametrize(
    "values",
    [
        {},
        {"template": "tea_purchase", "content": "不能同时提供两个正文来源"},
        {"template": "tea_purchase", "format": "txt"},
        {"template": "../../private"},
        {
            "template": "tea_purchase",
            "replacements": [{"part": "word/document.xml", "old_text": "甲方", "new_text": "乙方"}],
        },
        {"content": "正文", "replacements": [{"old_text": "正文", "new_text": "其他"}]},
    ],
)
def test_contract_creation_accepts_one_valid_content_source(values):
    with pytest.raises(ValueError):
        FileCreateInput.model_validate({"filename": "合同.docx", "format": "docx", **values})
