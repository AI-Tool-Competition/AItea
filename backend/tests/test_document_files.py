import base64
import io
import json
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.contexts.documents.infrastructure.document_files import create, edit, inspect, operate


TEXT = (
    "# 茶叶采购说明\n数量为100公斤，交货日期为[待补充]。\n\n| 产品 | 数量 |\n| --- | --- |\n| 六堡茶 | 100 |"
)


@pytest.mark.parametrize("format", ["docx", "pdf", "md", "txt", "csv", "xlsx", "pptx"])
def test_create_reopens_real_file_in_bounded_worker(format):
    content = TEXT
    if format == "xlsx":
        content = json.dumps({"sheets": [{"name": "报价", "rows": [["编号", "单价"], ["001", 120]]}]})
    elif format == "pptx":
        content = json.dumps({"slides": [{"title": "出口准备", "bullets": ["确认进口商", "保留原件"]}]})
    result = operate("create", {"format": format, "content": content})
    data = base64.b64decode(result["file_base64"])
    assert result["validation"]["reopened"]
    assert result["validation"]["visual_check"] == "not_performed"
    assert inspect(data, format)["total"] > 0
    if format == "docx":
        doc = Document(io.BytesIO(data))
        assert doc.tables[0].cell(1, 0).text == "六堡茶"
        assert "[待补充]" in doc.paragraphs[1].text
    if format == "pdf":
        reader = PdfReader(io.BytesIO(data))
        assert "茶叶采购说明" in reader.pages[0].extract_text()
    if format == "xlsx":
        book = load_workbook(io.BytesIO(data))
        assert book["报价"]["A2"].value == "001"
        book.close()


def package(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def test_word_cross_run_edit_preserves_styles_headers_tables_and_other_parts():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("支付货款").bold = True
    paragraph.add_run("100元，其他条件不变。").italic = True
    document.sections[0].header.paragraphs[0].text = "茶商原件"
    document.add_table(rows=1, cols=1).cell(0, 0).text = "不可更改"
    out = io.BytesIO()
    document.save(out)
    source = out.getvalue()
    result = edit(
        source,
        "docx",
        [
            {
                "part": "word/document.xml",
                "old_text": "货款100元",
                "new_text": "货款120元",
                "expected_count": 1,
            }
        ],
    )
    before, after = package(source), package(result)
    assert all(after[k] == v for k, v in before.items() if k != "word/document.xml")
    doc = Document(io.BytesIO(result))
    assert doc.paragraphs[0].text == "支付货款120元，其他条件不变。"
    assert doc.paragraphs[0].runs[0].bold
    assert doc.paragraphs[0].runs[1].italic
    assert doc.tables[0].cell(0, 0).text == "不可更改"
    assert doc.sections[0].header.paragraphs[0].text == "茶商原件"


def test_ambiguous_word_change_fails_without_saving():
    source = create("docx", "原金额100元\n另一条金额100元")
    with pytest.raises(ValueError, match="命中次数"):
        edit(
            source,
            "docx",
            [{"part": "word/document.xml", "old_text": "100元", "new_text": "120元", "expected_count": 1}],
        )


def test_spreadsheet_edit_preserves_formula_style_and_all_other_package_members():
    book = Workbook()
    sheet = book.active
    sheet.title = "报价"
    sheet.append(["001", 120, "=B1*2"])
    sheet["B1"].number_format = "0.00"
    sheet.merge_cells("A3:B3")
    sheet["A3"] = "合并标题"
    out = io.BytesIO()
    book.save(out)
    source = out.getvalue()
    result = edit(
        source, "xlsx", cells=[{"sheet": "报价", "cell": "B1", "expected_value": 120, "value": 125}]
    )
    before, after = package(source), package(result)
    assert all(after[k] == v for k, v in before.items() if k != "xl/worksheets/sheet1.xml")
    reopened = load_workbook(io.BytesIO(result), data_only=False)
    assert reopened["报价"]["A1"].value == "001"
    assert reopened["报价"]["B1"].value == 125
    assert reopened["报价"]["B1"].number_format == "0.00"
    assert reopened["报价"]["C1"].value == "=B1*2"
    assert "A3:B3" in reopened["报价"].merged_cells
    reopened.close()
    with pytest.raises(ValueError, match="非公式"):
        edit(source, "xlsx", cells=[{"sheet": "报价", "cell": "C1", "expected_value": None, "value": 250}])


def test_spreadsheet_literals_cannot_become_formulas_and_old_value_must_match():
    source = create(
        "xlsx",
        json.dumps({"sheets": [{"name": "报价", "rows": [["001", '=HYPERLINK("https://example.com")']]}]}),
    )
    book = load_workbook(io.BytesIO(source))
    assert book["报价"]["B1"].data_type == "s"
    book.close()
    with pytest.raises(ValueError, match="原值不一致"):
        edit(source, "xlsx", cells=[{"sheet": "报价", "cell": "A1", "expected_value": 1, "value": "002"}])


def test_ppt_edit_changes_only_requested_slide():
    source = create(
        "pptx",
        json.dumps(
            {
                "slides": [
                    {"title": "出口准备", "bullets": ["确认进口商"]},
                    {"title": "保留页面", "bullets": []},
                ]
            }
        ),
    )
    result = edit(
        source,
        "pptx",
        [
            {
                "part": "ppt/slides/slide1.xml",
                "old_text": "确认进口商",
                "new_text": "确认进口商资质",
                "expected_count": 1,
            }
        ],
    )
    before, after = package(source), package(result)
    assert all(after[k] == v for k, v in before.items() if k != "ppt/slides/slide1.xml")
    assert Presentation(io.BytesIO(result)).slides[0].placeholders[1].text == "确认进口商资质"


def test_text_edit_keeps_bom_leading_zeroes_and_line_endings():
    source = b"\xef\xbb\xbf" + "编号,价格\r\n001,120\r\n".encode()
    result = edit(source, "csv", [{"old_text": "001,120", "new_text": "001,125", "expected_count": 1}])
    assert result == b"\xef\xbb\xbf" + "编号,价格\r\n001,125\r\n".encode()


def test_unsupported_and_malformed_files_are_not_rewritten():
    with pytest.raises(ValueError, match="PDF"):
        edit(create("pdf", "原件"), "pdf", [{"old_text": "原件", "new_text": "副本", "expected_count": 1}])
    with pytest.raises(ValueError, match="列数"):
        create("docx", "| A | B |\n| --- | --- |\n| 1 |")
    with pytest.raises(ValueError):
        operate("inspect", {"format": "docx"}, b"not-office")
