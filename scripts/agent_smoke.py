"""Opt-in live-model smoke tests against a local running tea workspace.

Uses synthetic facts, isolated demo accounts/cases and actual tools. Reports
contain answers and bounded execution metadata, never provider reasoning or
credentials. Passing delivery checks does not certify legal interpretation.
"""

import argparse
import hashlib
import io
import json
import os
import re
import time
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from urllib.parse import urlsplit

import httpx
from docx import Document
from pypdf import PdfReader

LETTER = "催款函\n王芳：\n六堡茶货款8000元已逾期5天，请于本周内支付；如有困难，请先沟通。\n李明"
CONTRACT = """茶叶采购合同
买方：甲茶坊。卖方：乙茶厂。适用中国法。
1.甲向乙购买六堡茶100公斤，总价20000元。甲签约后三日内支付全部价款，乙以收到款项为条件安排生产。
2.交货时间为生产完成后，具体日期由乙通知。乙可委托其他工厂生产和交货；工厂未交货或产品有问题时，乙不承担任何责任。
3.甲须在收货当日验收，未在当日书面提出异议即视为全部质量合格。农残等需要送检才能发现的问题也适用同一期限。
4.检测标准和批次检测报告由乙自行决定是否提供。宣传资料中说产品适合出口，但本合同未约定具体出口国家。
5.乙延迟交货不承担责任，甲逾期付款每天支付合同总价5%的违约金。甲不得以质量异议为由要求协商退款。
6.交易发生争议时，甲只能到乙指定地点协商，不得向任何法院起诉。
7.本合同附件为产品清单，附件尚未填写批次、等级、包装和交货地址。
8.合同签字后生效。本段未写明的事实均未确认。
"""
SCENARIOS = [
    (
        "website",
        "请找广州茶博会主办方的官方网站，给出网站名称和可点击的完整网址，优先官方入口。",
        {"web_search"},
    ),
    (
        "business",
        "我是梧州六堡茶商，第一次找海外客户，预算两万元。比较与当地进口经销商合作和自营跨境网店两种方式，给推荐及先做的三件事。目的国尚未确定，不用联网，也不要先列办证清单。",
        set(),
    ),
    (
        "law_lookup",
        "只查中国法：《中华人民共和国民法典》第五百七十九条的完整原文和官方出处。不要扩展其他法域。",
        {"lookup_chinese_law", "read_law"},
    ),
    (
        "payment",
        "假设中国国内茶叶买卖合同有效，我作为卖方已按约交货，买方收货后付款期限已过20天，无质量争议、合同变更或其他正当抗辩。可以要求货款及违约责任吗？合同未约定违约金，说明法律依据及下一步。不要重复询问题目已明确的事实。",
        {"lookup_chinese_law", "read_law", "search_legal_evidence"},
    ),
    (
        "contract",
        "请站在买方甲茶坊立场审查以下合同，优先指出最影响付款、交货和质量追责的三项风险，并给可直接替换的三段文字。只交付审查意见，不保存文件；需要法律依据时查原文。\n"
        + CONTRACT,
        set(),
    ),
    (
        "files",
        "请按下列中文正文原样生成Word和同内容PDF，保留姓名和金额，不补写内容：\n" + LETTER,
        {"create_document_file"},
    ),
    (
        "missing_facts",
        "准备给海外客户寄六堡茶样品，目的国家和进口商都还没有确定。先说现在能整理的产品和交易资料，再只问最影响下一步的一项信息。不要默认目的国是业务模板里的国家，不用联网。",
        set(),
    ),
    (
        "business_card",
        "我叫李明，在广西梧州经营六堡茶，现有六堡茶200公斤，准备先做海外批发。请把已经说清楚的经营者和货源信息填写到右侧业务卡，未知项留空，再告诉我下一步先确认什么。",
        {"update_business_card"},
    ),
]


def checked(response):
    response.raise_for_status()
    return response.json()


def perform(client, case, conversation, label, question, required):
    started = time.monotonic()
    record = {"scenario": label, "case_id": case["id"], "checks": {}}
    blobs = []
    try:
        response = client.post(
            f"/api/tea/cases/{case['id']}/messages",
            json={
                "stage_id": case["stages"][0]["id"],
                "conversation_id": conversation["id"],
                "content": question,
                "client_message_id": str(uuid.uuid4()),
            },
        )
        record["http_status"] = response.status_code
        payload = checked(response)
        answer = payload["messages"][-1]
        record["answer"] = answer["content"]
        record["citations"] = answer.get("citations", [])
        record["delivery_status"] = next(
            (s["delivery_status"] for s in answer.get("trace", []) if s.get("delivery_status")), "unknown"
        )
        record["tools"] = [
            {k: t[k] for k in ("name", "status", "attempt", "duration_ms", "error_code") if k in t}
            for step in answer.get("trace", [])
            for t in step.get("tools", [])
        ]
        succeeded = {t["name"] for t in record["tools"] if t["status"] == "completed"}
        record["checks"] = {
            "answer_present": bool(record["answer"].strip()),
            "delivery_complete": record["delivery_status"] == "answered",
            "required_tool_executed": not required or bool(succeeded & required),
            "no_internal_stage_names": not bool(
                re.search(r"QA-0[1-6]|pending_commit|附条件结论·待复核", record["answer"])
            ),
            "no_unrecovered_tool_failure": not any(
                t["status"] in {"failed", "cancelled"} for t in record["tools"]
            ),
        }
        if label in {"law_lookup", "payment"}:
            record["checks"]["citations_present"] = bool(record["citations"])
        if label == "website":
            record["checks"]["website_link_present"] = bool(re.search(r"https?://[^\s)]+", record["answer"]))
        if label == "business_card":
            current = checked(client.get(f"/api/tea/cases/{case['id']}"))
            values = json.dumps(
                [task["values"] for stage in current["stages"] for task in stage["tasks"]],
                ensure_ascii=False,
            )
            record["checks"]["facts_actually_saved"] = "李明" in values and "200" in values
        record["downloads"] = []
        for url in dict.fromkeys(re.findall(r"\]\((/api/tea/documents/[^)]+/download)\)", record["answer"])):
            file = client.get(url)
            record["downloads"].append({"url": url, "status": file.status_code, "size": len(file.content)})
            if file.is_success:
                blobs.append(file.content)
    except Exception as error:  # noqa: BLE001 - Record an isolated scenario failure and continue.
        record["error_type"] = type(error).__name__
        record["checks"]["request_completed"] = False
    record["seconds"] = round(time.monotonic() - started, 3)
    record["within_90_seconds"] = record["seconds"] <= 90
    return record, blobs


def run_group(base, round_number, scenario):
    label, question, required = scenario
    with httpx.Client(base_url=base, timeout=httpx.Timeout(210, connect=10), trust_env=False) as client:
        checked(client.post("/api/tea/auth/demo", json={}))
        template = checked(client.get("/api/tea/bootstrap"))["templates"][0]
        mode = next(m for m in template["modes"] if m["enabled"])
        case = checked(
            client.post(
                "/api/tea/cases",
                json={
                    "title": f"稳定性验证 {label} {round_number}",
                    "company": "",
                    "product": "六堡茶",
                    "template_id": template["id"],
                    "mode_id": mode["id"],
                },
            )
        )
        conversation = checked(
            client.post(
                f"/api/tea/cases/{case['id']}/stages/{case['stages'][0]['id']}/conversations",
                json={"client_request_id": str(uuid.uuid4()), "title": label},
            )
        )
        first, blobs = perform(client, case, conversation, label, question, required)
        first["round"] = round_number
        records = [first]
        if label == "files":
            word = next((b for b in blobs if b.startswith(b"PK")), None)
            pdf = next((b for b in blobs if b.startswith(b"%PDF-")), None)
            original = "\n".join(p.text for p in Document(io.BytesIO(word)).paragraphs) if word else ""
            first["checks"]["word_and_pdf_downloaded"] = bool(word and pdf)
            first["checks"]["word_content_preserved"] = original == LETTER
            first["checks"]["pdf_content_present"] = bool(pdf) and "8000" in "\n".join(
                p.extract_text() or "" for p in PdfReader(io.BytesIO(pdf)).pages
            )
            if word:
                edited, files = perform(
                    client,
                    case,
                    conversation,
                    "file_edit",
                    "将刚生成的Word中8000元改为8500元，其他文字和格式不变。先读取原件再定点修改，另存Word副本，原件保留，不要新写全文，也不需要PDF。",
                    {"edit_document_file"},
                )
                edited["round"] = round_number
                new_word = next((b for b in files if b.startswith(b"PK")), None)
                new_text = (
                    "\n".join(p.text for p in Document(io.BytesIO(new_word)).paragraphs) if new_word else ""
                )
                edited["checks"]["only_requested_text_changed"] = bool(
                    new_word
                ) and new_text == original.replace("8000", "8500")
                original_url = next(
                    d["url"] for d in first["downloads"] if client.get(d["url"]).content.startswith(b"PK")
                )
                edited["checks"]["original_unchanged"] = (
                    hashlib.sha256(client.get(original_url).content).digest() == hashlib.sha256(word).digest()
                )
                records.append(edited)
        for record in records:
            record["delivery_pass"] = bool(record["checks"]) and all(record["checks"].values())
        return records


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base", default="http://127.0.0.1:8483")
    parser.add_argument("--report", required=True)
    parser.add_argument("--rounds", type=int, default=2, choices=(1, 2, 3))
    parser.add_argument("--jobs", type=int, default=2, choices=(1, 2))
    parser.add_argument("--scenarios", nargs="+", choices=[s[0] for s in SCENARIOS])
    args = parser.parse_args()
    selected = [s for s in SCENARIOS if not args.scenarios or s[0] in args.scenarios]
    if urlsplit(args.base).hostname not in {"127.0.0.1", "localhost", "::1"}:
        parser.error("仅允许本地隔离验证实例")
    os.umask(0o077)
    target = Path(args.report)
    target.parent.mkdir(parents=True, exist_ok=True)
    report = {
        "base": args.base,
        "parallelism": args.jobs,
        "rounds": args.rounds,
        "runs": [],
        "note": "Real models/tools. Delivery checks do not certify legal correctness.",
    }
    with ThreadPoolExecutor(max_workers=args.jobs) as pool:
        futures = {
            pool.submit(run_group, args.base, n, scenario): (n, scenario[0])
            for n in range(1, args.rounds + 1)
            for scenario in selected
        }
        for future in as_completed(futures):
            try:
                records = future.result()
            except Exception as error:  # noqa: BLE001 - Record an isolated scenario failure and continue.
                n, label = futures[future]
                records = [
                    {
                        "round": n,
                        "scenario": label,
                        "delivery_pass": False,
                        "error_type": type(error).__name__,
                    }
                ]
            report["runs"].extend(records)
            target.write_text(json.dumps(report, ensure_ascii=False, indent=2))
            for record in records:
                print(
                    json.dumps(
                        {
                            k: record[k]
                            for k in (
                                "round",
                                "scenario",
                                "seconds",
                                "delivery_pass",
                                "checks",
                            )
                            if k in record
                        },
                        ensure_ascii=False,
                    ),
                    flush=True,
                )
    passed = sum(r["delivery_pass"] for r in report["runs"])
    print(
        json.dumps(
            {
                "runs": len(report["runs"]),
                "delivery_passed": passed,
                "report": str(target),
            },
            ensure_ascii=False,
        )
    )
    expected = args.rounds * (len(selected) + int(any(s[0] == "files" for s in selected)))
    return 0 if passed == expected else 1


if __name__ == "__main__":
    raise SystemExit(main())
